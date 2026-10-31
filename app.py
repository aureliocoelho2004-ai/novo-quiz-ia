"""Quizzify - transforma PDFs, Word, PowerPoint e fotos em quizzes/flashcards com IA.

Toda a lógica fica aqui; o HTML vive em templates/ e o JS/CSS da página em static/.
"""
import hashlib
import hmac
import io
import json
import logging
import os
import re
import threading
import unicodedata
import time
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from functools import wraps

import requests
from flask import (Flask, Response, abort, flash, g, has_request_context, jsonify, redirect,
                   render_template, request, session, url_for)
from itsdangerous import BadSignature, URLSafeTimedSerializer
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.security import check_password_hash, generate_password_hash

log = logging.getLogger("quizify")
logging.basicConfig(level=logging.INFO)

MAX_ARQUIVOS = 5
MAX_TOTAL_MB = 50
MIN_ITENS, MAX_ITENS = 5, 50
MAX_CARACTERES_TEXTO = 60_000
LIMITE_DIARIO = int(os.environ.get("LIMITE_DIARIO", 50))           # gerações por utilizador por dia
LIMITE_TESTE = int(os.environ.get("LIMITE_TESTE", 10))             # gerações por dia em contas de teste
DURACOES_TESTE = {"24h": 24, "3d": 72, "7d": 168}                  # opções do painel, em horas
INTERVALO_SEGUNDOS = int(os.environ.get("INTERVALO_SEGUNDOS", 30))  # espera entre gerações
FUSO_BRASIL = timezone(timedelta(hours=-3))
DIA = 24 * 3600

# Assinaturas pagas pela InfinitePay (Checkout Integrado). Preços em centavos.
PLANOS = {
    "mensal": {"nome": "Mensal", "centavos": int(os.environ.get("PRECO_MENSAL_CENTAVOS", 2990)), "dias": 30},
    "anual": {"nome": "Anual", "centavos": int(os.environ.get("PRECO_ANUAL_CENTAVOS", 19700)), "dias": 365},
}
INFINITEPAY_HANDLE = os.environ.get("INFINITEPAY_HANDLE", "").strip().lstrip("$@")
INFINITEPAY_API = "https://api.checkout.infinitepay.io"
AVISO_VENCIMENTO_DIAS = 5
EXTENSOES_DOC = {".pdf", ".docx", ".pptx"}
EXTENSOES_IMG = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}
MIME_IMG = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
            ".webp": "image/webp", ".heic": "image/heic", ".heif": "image/heif"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_TOTAL_MB * 1024 * 1024
app.json.ensure_ascii = False
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("RENDER")),  # Render serve sempre por HTTPS
    PERMANENT_SESSION_LIFETIME=60 * 60 * 24 * 7,
)


def _chave_secreta():
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    # Sem SECRET_KEY, deriva uma chave estável de um segredo existente para que
    # todos os workers do gunicorn aceitem a mesma sessão.
    base = os.environ.get("FIREBASE_KEY") or os.environ.get("GEMINI_API_KEY")
    if base:
        log.warning("SECRET_KEY não definida; usando chave derivada. Defina SECRET_KEY no Render.")
        return hashlib.sha256(("quizify-sessao:" + base).encode()).hexdigest()
    log.warning("SECRET_KEY não definida; sessões serão perdidas ao reiniciar.")
    return os.urandom(32).hex()


app.secret_key = _chave_secreta()


# --------------------------------------------------------------------------- #
# Serviços externos (inicializados sob demanda para o app subir mesmo sem
# credenciais e o /health continuar a responder ao cron-job).
# --------------------------------------------------------------------------- #
_cliente_gemini = None
_bucket = None
# O gunicorn roda com threads (gunicorn.conf.py): várias requisições ao mesmo
# tempo no mesmo processo, por isso a inicialização e os contadores usam travas.
_trava_init = threading.Lock()
_trava_contadores = threading.Lock()


def obter_segredo(chave):
    valor = os.environ.get(chave)
    if not valor:
        raise RuntimeError(f"Variável de ambiente ausente: {chave}")
    return valor


def gemini():
    global _cliente_gemini
    if _cliente_gemini is None:
        with _trava_init:
            if _cliente_gemini is None:
                from google import genai
                _cliente_gemini = genai.Client(api_key=obter_segredo("GEMINI_API_KEY"))
    return _cliente_gemini


def bucket_firebase():
    global _bucket
    if _bucket is None:
        with _trava_init:
            if _bucket is None:
                import firebase_admin
                from firebase_admin import credentials, storage
                if not firebase_admin._apps:
                    cred = credentials.Certificate(json.loads(obter_segredo("FIREBASE_KEY")))
                    firebase_admin.initialize_app(cred, {"storageBucket": obter_segredo("FIREBASE_BUCKET")})
                _bucket = storage.bucket()
    return _bucket


# Gerações simultâneas por processo: cada uma lê arquivos e imagens na memória,
# então limitamos para não estourar a RAM do Render. Quem passar do limite
# espera a vez (até ESPERA_VAGA_SEGUNDOS) em vez de derrubar o servidor.
GERACOES_SIMULTANEAS = int(os.environ.get("GERACOES_SIMULTANEAS", 4))
ESPERA_VAGA_SEGUNDOS = int(os.environ.get("ESPERA_VAGA_SEGUNDOS", 90))
_vagas_geracao = threading.BoundedSemaphore(GERACOES_SIMULTANEAS)


@contextmanager
def vaga_geracao():
    conseguiu = _vagas_geracao.acquire(timeout=ESPERA_VAGA_SEGUNDOS)
    try:
        yield conseguiu
    finally:
        if conseguiu:
            _vagas_geracao.release()


# --------------------------------------------------------------------------- #
# Utilizadores e autenticação
# Guardados em usuarios.json no Firebase, com senhas em hash:
#   {"nome": {"senha": "<hash>", "admin": bool, "trocar_senha": bool}}
# --------------------------------------------------------------------------- #
ARQUIVO_USUARIOS = "usuarios.json"
MAX_TENTATIVAS, JANELA_TENTATIVAS = 5, 10 * 60
_tentativas = {}


class FalhaUsuarios(Exception):
    """Não foi possível ler a lista de utilizadores do Firebase."""


def carregar_usuarios(estrito=False):
    """Com estrito=True, uma falha de leitura levanta FalhaUsuarios em vez de
    devolver uma lista vazia — obrigatório antes de gravar, senão a gravação
    apagaria todos os utilizadores existentes."""
    try:
        dados, geracao = ler_json_versionado(ARQUIVO_USUARIOS)
        dados = dados or {}
    except Exception as erro:
        log.exception("Falha ao ler utilizadores do Firebase")
        if estrito:
            raise FalhaUsuarios from erro
        dados, geracao = {}, None
    if has_request_context():
        g.geracao_usuarios = geracao

    usuarios = {}
    for nome, valor in dados.items():
        if isinstance(valor, str):  # formato antigo (Streamlit): senha em texto puro
            valor = {"senha": generate_password_hash(valor), "admin": nome == "admin", "trocar_senha": True}
        usuarios[nome] = valor

    # Primeiro acesso: cria o admin a partir das variáveis de ambiente.
    if not usuarios and os.environ.get("ADMIN_PASSWORD"):
        nome = os.environ.get("ADMIN_USER", "admin").strip().lower()
        usuarios[nome] = {"senha": generate_password_hash(os.environ["ADMIN_PASSWORD"]),
                          "admin": True, "trocar_senha": False}
    return usuarios


def salvar_usuarios(usuarios):
    """Grava só se ninguém alterou o arquivo desde a leitura (senão levanta
    PreconditionFailed), para um pagamento e o admin não se sobrescreverem."""
    geracao = g.get("geracao_usuarios") if has_request_context() else None
    g.geracao_usuarios = gravar_json_versionado(ARQUIVO_USUARIOS, usuarios, geracao)


def alterar_usuarios(alteracao, tentativas=5):
    """Lê, aplica `alteracao(usuarios)` e grava, repetindo se outra gravação
    acontecer no meio. A alteração devolve False para cancelar sem gravar."""
    from google.api_core.exceptions import PreconditionFailed
    for _ in range(tentativas):
        usuarios = carregar_usuarios(estrito=True)
        resultado = alteracao(usuarios)
        if resultado is False:
            return False
        try:
            salvar_usuarios(usuarios)
            return resultado
        except PreconditionFailed:
            continue
    raise FalhaUsuarios("Muitas gravações simultâneas")


def ler_json_versionado(caminho):
    """Devolve (dados, geração); geração 0 significa que o arquivo não existe."""
    blob = bucket_firebase().get_blob(caminho)
    if blob is None:
        return None, 0
    return json.loads(blob.download_as_bytes(if_generation_match=blob.generation)), blob.generation


def gravar_json_versionado(caminho, dados, geracao=None):
    blob = bucket_firebase().blob(caminho)
    blob.upload_from_string(json.dumps(dados, ensure_ascii=False), content_type="application/json",
                            if_generation_match=geracao)
    return blob.generation


def ip_cliente():
    return (request.headers.get("X-Forwarded-For") or request.remote_addr or "").split(",")[0].strip()


def bloqueado(ip):
    agora = time.time()
    with _trava_contadores:
        _tentativas[ip] = [t for t in _tentativas.get(ip, []) if agora - t < JANELA_TENTATIVAS]
        return len(_tentativas[ip]) >= MAX_TENTATIVAS


def registrar_tentativa(contadores, ip, quando=None):
    with _trava_contadores:
        contadores.setdefault(ip, []).append(quando or time.time())


def senha_valida(senha):
    return len(senha) >= 8


def expirou(conta):
    """Contas de teste guardam "expira_em" (timestamp); contas normais não expiram."""
    return bool(conta.get("expira_em")) and time.time() >= conta["expira_em"]


def email_da_conta(nome, conta):
    """E-mail para avisos: o campo "email" ou o próprio nome, se for um e-mail."""
    return conta.get("email") or (nome if "@" in nome else None)


def converter_teste_expirado(nome):
    """Teste acabou: a conta vira assinante sem pagamento (mantém os quizzes),
    e a pessoa cai na página dos planos em vez de ficar bloqueada."""
    def converter(usuarios):
        conta = usuarios.get(nome)
        if not conta:
            return False
        if conta.get("expira_em") and time.time() >= conta["expira_em"]:
            conta.pop("expira_em")
            conta.update(assinante=True, pago_ate=conta.get("pago_ate") or 0, teste_acabou_em=time.time())
            conta.setdefault("origem", "teste")
        return conta
    return alterar_usuarios(converter)


def entrar_na_sessao(nome, conta):
    session.clear()
    session.permanent = True
    session.update(usuario=nome, admin=bool(conta.get("admin")), trocar_senha=bool(conta.get("trocar_senha")),
                   expira_em=conta.get("expira_em"), tem_email=bool(email_da_conta(nome, conta)))
    atualizar_sessao_assinatura(conta)


def assinatura_ativa(conta):
    """Contas de assinante só funcionam com pagamento em dia; as outras sempre."""
    return not conta.get("assinante") or conta.get("pago_ate", 0) > time.time()


def atualizar_sessao_assinatura(conta):
    session.update(assinante=bool(conta.get("assinante")), pago_ate=conta.get("pago_ate", 0))


def conta_ainda_existe(usuarios):
    """Encerra a sessão se o utilizador foi removido ou o teste expirou depois do login."""
    conta = usuarios.get(session.get("usuario"))
    if conta and not expirou(conta):
        return True
    session.clear()
    return False


@app.template_filter("data_hora")
def data_hora(timestamp):
    return datetime.fromtimestamp(timestamp, FUSO_BRASIL).strftime("%d/%m às %H:%M")


def _blob_uso(usuario):
    return bucket_firebase().blob(f"uso/{usuario}.json")


def verificar_limite(usuario, limite):
    """Devolve uma mensagem de erro se o utilizador passou do limite, senão None.
    Guardado no Firebase para valer entre workers e reinícios do Render.
    Administradores não têm limite."""
    hoje = datetime.now(FUSO_BRASIL).date().isoformat()
    try:
        blob = _blob_uso(usuario)
        uso = json.loads(blob.download_as_bytes()) if blob.exists() else {}
    except Exception:
        log.exception("Falha ao ler uso de %s", usuario)
        return None  # não bloqueia o professor por uma falha de leitura
    if uso.get("dia") != hoje:
        uso = {"dia": hoje, "total": 0, "ultimo": 0}

    espera = INTERVALO_SEGUNDOS - (time.time() - uso.get("ultimo", 0))
    if espera > 0:
        return f"⏳ Aguarde {int(espera) + 1}s para criar outro conteúdo."
    if uso["total"] >= limite:
        return f"⚠️ Limite de {limite} conteúdos por dia atingido. Tente novamente amanhã."

    uso["ultimo"] = time.time()  # marca já, para impedir pedidos simultâneos
    _gravar_uso(usuario, uso)
    return None


def registrar_geracao(usuario):
    hoje = datetime.now(FUSO_BRASIL).date().isoformat()
    try:
        blob = _blob_uso(usuario)
        uso = json.loads(blob.download_as_bytes()) if blob.exists() else {}
    except Exception:
        log.exception("Falha ao ler uso de %s", usuario)
        return
    if uso.get("dia") != hoje:
        uso = {"dia": hoje, "total": 0, "ultimo": time.time()}
    uso["total"] = uso.get("total", 0) + 1
    _gravar_uso(usuario, uso)


def _gravar_uso(usuario, uso):
    try:
        _blob_uso(usuario).upload_from_string(json.dumps(uso), content_type="application/json")
    except Exception:
        log.exception("Falha ao gravar uso de %s", usuario)


def assinatura_paga_agora():
    """A sessão diz que venceu: confirma no Firebase antes de bloquear, porque o
    pagamento pode ter sido confirmado (aviso da InfinitePay) depois do login."""
    conta = carregar_usuarios().get(session.get("usuario"))
    if conta:
        atualizar_sessao_assinatura(conta)
    return bool(conta) and assinatura_ativa(conta)


# Com a assinatura vencida a pessoa só pode pagar, trocar a senha e sair. "Meus quizzes"
# fica bloqueado até renovar (o histórico não é apagado e volta com o pagamento).
ROTAS_SEM_ASSINATURA = {"assinatura", "assinatura_pagar", "assinatura_retorno", "trocar_senha", "logout", "meu_email"}


def login_obrigatorio(view):
    @wraps(view)
    def envolvida(*args, **kwargs):
        if not session.get("usuario"):
            if request.path.startswith("/api/"):
                return jsonify(erro="Sessão expirada. Entre novamente."), 401
            # full_path mantém os parâmetros (ex.: retorno do pagamento da InfinitePay)
            return redirect(url_for("login", proximo=request.full_path.rstrip("?")))
        if session.get("expira_em") and time.time() >= session["expira_em"]:
            if request.path.startswith("/api/"):
                return jsonify(erro="O seu teste acabou. Escolha um plano para continuar."), 402
            try:
                conta = converter_teste_expirado(session["usuario"])
            except Exception:
                log.exception("Falha ao converter teste de %s", session["usuario"])
                session.clear()
                return redirect(url_for("login"))
            if not conta:
                session.clear()
                return redirect(url_for("login"))
            entrar_na_sessao(session["usuario"], conta)
            if conta.get("expira_em") or not conta.get("assinante") or (conta.get("pago_ate") or 0) > time.time():
                return view(*args, **kwargs)      # teste prolongado, ou já assinou
            flash("O seu teste acabou. Escolha um plano para continuar — os seus quizzes estão guardados. 💾", "erro")
            return redirect(url_for("assinatura", plano="anual"))
        if session.get("trocar_senha") and request.endpoint not in ("trocar_senha", "logout"):
            return redirect(url_for("trocar_senha"))
        if (session.get("assinante") and session.get("pago_ate", 0) <= time.time()
                and request.endpoint not in ROTAS_SEM_ASSINATURA
                and not assinatura_paga_agora()):
            if request.path.startswith("/api/"):
                return jsonify(erro="A sua assinatura não está ativa. Renove para continuar."), 402
            return redirect(url_for("assinatura"))
        return view(*args, **kwargs)
    return envolvida


def admin_obrigatorio(view):
    @wraps(view)
    @login_obrigatorio
    def envolvida(*args, **kwargs):
        if not session.get("admin"):
            abort(403)
        return view(*args, **kwargs)
    return envolvida


# --------------------------------------------------------------------------- #
# Extração de conteúdo
# --------------------------------------------------------------------------- #
def extensao(nome):
    return os.path.splitext(nome or "")[1].lower()


def extrair_pdf(dados):
    from pypdf import PdfReader
    leitor = PdfReader(io.BytesIO(dados))
    return "\n".join(t for p in leitor.pages if (t := p.extract_text()))


def extrair_docx(dados):
    import docx
    documento = docx.Document(io.BytesIO(dados))
    linhas = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            linhas.append(" | ".join(c.text.strip() for c in linha.cells if c.text.strip()))
    return "\n".join(linhas)


def extrair_pptx(dados):
    import pptx
    linhas = []
    for slide in pptx.Presentation(io.BytesIO(dados)).slides:
        for forma in slide.shapes:
            if forma.has_text_frame:
                linhas.extend(p.text.strip() for p in forma.text_frame.paragraphs if p.text.strip())
    return "\n".join(linhas)


def reduzir_imagem(dados, mime):
    """Rede de segurança: o navegador já comprime, mas imagens grandes que
    escaparem (ex.: HEIC que o navegador não lê) são reduzidas aqui."""
    if len(dados) <= 1_500_000:
        return dados, mime
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(dados))
        img.thumbnail((1600, 1600))
        saida = io.BytesIO()
        img.convert("RGB").save(saida, "JPEG", quality=80)
        return saida.getvalue(), "image/jpeg"
    except Exception:
        return dados, mime


def preparar_conteudo(arquivos):
    """Devolve (texto_extraido, partes_binarias) para enviar ao Gemini numa só chamada."""
    from google.genai import types

    textos, partes = [], []
    for arq in arquivos:
        ext = extensao(arq.filename)
        dados = arq.read()
        if not dados:
            continue
        if ext == ".pdf":
            texto = extrair_pdf(dados)
            if texto.strip():
                textos.append(texto)
            else:  # PDF digitalizado (só imagens): o Gemini lê o PDF diretamente
                partes.append(types.Part.from_bytes(data=dados, mime_type="application/pdf"))
        elif ext == ".docx":
            textos.append(extrair_docx(dados))
        elif ext == ".pptx":
            textos.append(extrair_pptx(dados))
        elif ext in EXTENSOES_IMG or (arq.mimetype or "").startswith("image/"):
            # Chrome/Windows envia HEIC como "application/octet-stream", que a IA recusa:
            # nesse caso o tipo vem da extensão.
            mime = arq.mimetype if (arq.mimetype or "").startswith("image/") else MIME_IMG.get(ext, "image/jpeg")
            dados, mime = reduzir_imagem(dados, mime)
            partes.append(types.Part.from_bytes(data=dados, mime_type=mime))
    return "\n\n".join(t for t in textos if t.strip())[:MAX_CARACTERES_TEXTO], partes


# --------------------------------------------------------------------------- #
# Geração com IA
# --------------------------------------------------------------------------- #
PROMPT_QUIZ = """Você é um professor experiente. Com base EXCLUSIVAMENTE no material fornecido
(texto e/ou imagens), crie {n} perguntas de múltipla escolha em português do Brasil.
Regras:
- Cada pergunta tem exatamente 4 opções, apenas UMA correta.
- Varie a posição da resposta correta.
- "correta" é o ÍNDICE (0 a 3) da opção correta.
- "explicacao" justifica a resposta em 1 ou 2 frases.
Responda apenas com um array JSON no formato:
[{{"pergunta": "...", "opcoes": ["...", "...", "...", "..."], "correta": 0, "explicacao": "..."}}]"""

PROMPT_FLASH = """Você é um professor experiente. Com base EXCLUSIVAMENTE no material fornecido
(texto e/ou imagens), crie {n} flashcards em português do Brasil com os conceitos mais importantes.
Responda apenas com um array JSON no formato:
[{{"frente": "conceito ou pergunta curta", "verso": "explicação ou resposta"}}]"""


def chamar_gemini(prompt, texto, partes):
    from google.genai import types

    conteudo = [prompt]
    if texto:
        conteudo.append(f"MATERIAL:\n{texto}")
    conteudo.extend(partes)
    resposta = gemini().models.generate_content(
        model=os.environ.get("GEMINI_MODEL", "gemini-3.8-flash"),
        contents=conteudo,
        config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.4),
    )
    bruto = re.sub(r"^```(?:json)?\s*|\s*```$", "", (resposta.text or "").strip())
    dados = json.loads(bruto)
    if isinstance(dados, dict):  # às vezes vem embrulhado: {"perguntas": [...]}
        dados = next((v for v in dados.values() if isinstance(v, list)), [])
    return dados


def normalizar_quiz(itens, n):
    quiz = []
    for item in itens:
        if not isinstance(item, dict) or not isinstance(item.get("opcoes"), list):
            continue  # um item mal formado não derruba o quiz inteiro
        try:
            correta = int(item.get("correta", item.get("resposta_correta", -1)))
        except (TypeError, ValueError):
            continue
        # Remove opções vazias recalculando o índice da correta (antes ele deslocava
        # e o quiz marcava como certa a opção errada).
        validas = [(i, str(o).strip()) for i, o in enumerate(item["opcoes"]) if str(o).strip()]
        novos = [i for i, _ in validas]
        opcoes = [texto for _, texto in validas]
        if not item.get("pergunta") or len(opcoes) < 2 or correta not in novos:
            continue
        quiz.append({"id": len(quiz), "pergunta": str(item["pergunta"]).strip(), "opcoes": opcoes,
                     "correta": novos.index(correta), "explicacao": str(item.get("explicacao", "")).strip()})
    return quiz[:n]


def normalizar_flash(itens, n):
    return [{"frente": str(i["frente"]).strip(), "verso": str(i["verso"]).strip()}
            for i in itens if isinstance(i, dict) and i.get("frente") and i.get("verso")][:n]


# --------------------------------------------------------------------------- #
# Publicação
# --------------------------------------------------------------------------- #
def nome_seguro(texto):
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-zA-Z0-9]+", "-", texto).strip("-").lower()[:40] or "quizzify"


def publicar_firebase(conteudo_html, prefixo):
    """Devolve (link público, caminho no bucket), ou (None, None) se falhar."""
    try:
        caminho = f"quizzes/{prefixo}_{uuid.uuid4().hex[:10]}.html"
        blob = bucket_firebase().blob(caminho)
        blob.cache_control = "public, max-age=3600"
        # charset explícito: sem ele acentos/emojis aparecem corrompidos no Android/iOS
        blob.upload_from_string(conteudo_html.encode("utf-8"), content_type="text/html; charset=utf-8")
        blob.make_public()
        return blob.public_url, caminho
    except Exception:
        log.exception("Falha ao publicar no Firebase")
        return None, None


# --------------------------------------------------------------------------- #
# Histórico "Meus quizzes": uma lista por utilizador em historico/<hash>.json,
# do mais recente para o mais antigo. Remover da lista não apaga o quiz, para
# os links já enviados aos alunos continuarem a funcionar.
# --------------------------------------------------------------------------- #
HISTORICO_MAX = 300


def _caminho_historico(usuario):
    return f"historico/{hashlib.sha256(usuario.encode()).hexdigest()[:32]}.json"


def listar_historico(usuario):
    try:
        itens, _ = ler_json_versionado(_caminho_historico(usuario))
        return itens or []
    except Exception:
        log.exception("Falha ao ler histórico de %s", usuario)
        return []


def alterar_historico(usuario, alteracao, tentativas=5):
    from google.api_core.exceptions import PreconditionFailed
    caminho = _caminho_historico(usuario)
    for _ in range(tentativas):
        itens, geracao = ler_json_versionado(caminho)
        itens = itens or []
        if alteracao(itens) is False:
            return False
        try:
            gravar_json_versionado(caminho, itens, geracao)
            return True
        except PreconditionFailed:
            continue
    raise RuntimeError("Muitas gravações simultâneas no histórico")


def registrar_no_historico(usuario, item):
    def acrescentar(itens):
        itens.insert(0, item)
        del itens[HISTORICO_MAX:]
    try:
        alterar_historico(usuario, acrescentar)
    except Exception:
        log.exception("Falha ao gravar histórico de %s", usuario)  # o quiz já foi entregue; só não entra na lista


# --------------------------------------------------------------------------- #
# Pagamentos (InfinitePay — Checkout Integrado)
# Cada compra gera um pedido em pedidos/<order_nsu>.json. O aviso da
# InfinitePay não é assinado, por isso o pagamento é SEMPRE confirmado na
# própria InfinitePay (payment_check) antes de liberar o acesso.
# --------------------------------------------------------------------------- #
class FalhaPagamento(Exception):
    """A InfinitePay não respondeu como esperado."""


def url_publica(endpoint, **valores):
    base = os.environ.get("URL_PUBLICA", "").rstrip("/") or request.host_url.rstrip("/").replace("http://", "https://")
    return base + url_for(endpoint, **valores)


def url_webhook(endpoint, **valores):
    """Endereço para avisos de pagamento. Usa o endereço fixo do Render
    (RENDER_EXTERNAL_URL, ex.: https://quizify.onrender.com), que não depende do
    DNS do domínio próprio: se o DNS falhar, os pagamentos continuam a ser creditados."""
    base = (os.environ.get("URL_WEBHOOK") or os.environ.get("RENDER_EXTERNAL_URL") or "").rstrip("/")
    return base + url_for(endpoint, **valores) if base else url_publica(endpoint, **valores)


def token_webhook():
    """Segredo incluído no endereço do aviso, para descartar chamadas de terceiros."""
    return os.environ.get("INFINITEPAY_WEBHOOK_TOKEN") or hmac.new(
        app.secret_key.encode(), b"infinitepay-webhook", hashlib.sha256).hexdigest()[:32]


def _chamar_infinitepay(caminho, corpo):
    cabecalhos = {"Content-Type": "application/json", "Accept": "application/json"}
    if os.environ.get("INFINITEPAY_API_KEY"):
        cabecalhos["X-Api-Key"] = os.environ["INFINITEPAY_API_KEY"]
    try:
        resposta = requests.post(INFINITEPAY_API + caminho, json=corpo, headers=cabecalhos, timeout=15)
        dados = resposta.json()
    except (requests.RequestException, ValueError) as erro:
        raise FalhaPagamento(f"Sem resposta válida da InfinitePay em {caminho}") from erro
    if not resposta.ok:
        raise FalhaPagamento(f"InfinitePay {caminho} devolveu {resposta.status_code}: {dados}")
    return dados


def criar_checkout(pedido_id, plano, email):
    dados = _chamar_infinitepay("/links", {
        "handle": INFINITEPAY_HANDLE,
        "order_nsu": pedido_id,
        "redirect_url": url_publica("assinatura_retorno"),
        "webhook_url": url_webhook("webhook_infinitepay", token=token_webhook()),
        "items": [{"quantity": 1, "price": plano["centavos"], "description": f"Quizzify - Plano {plano['nome']}"}],
        "customer": {"email": email} if "@" in email else {},
    })
    link = next((dados.get(c) for c in ("url", "link", "checkout_url", "payment_url") if dados.get(c)), None)
    if not link:
        raise FalhaPagamento(f"InfinitePay não devolveu o link: {dados}")
    return link


def confirmar_pagamento(pedido_id, transacao, slug):
    """Confirma na InfinitePay e soma os dias à conta. Seguro para chamar várias
    vezes (aviso + retorno do cliente): o pedido só é creditado uma vez.
    Devolve "pago", "ja_pago", "nao_pago" ou "desconhecido"."""
    from google.api_core.exceptions import PreconditionFailed

    if not re.fullmatch(r"[0-9a-f]{32}", pedido_id):
        return "desconhecido"
    caminho = f"pedidos/{pedido_id}.json"
    pedido, geracao = ler_json_versionado(caminho)
    if not pedido:
        return "desconhecido"
    if pedido.get("status") == "pago":
        return "ja_pago"

    check = _chamar_infinitepay("/payment_check", {
        "handle": INFINITEPAY_HANDLE, "order_nsu": pedido_id, "transaction_nsu": transacao, "slug": slug})
    pago = check.get("paid") is True or str(check.get("paid")).lower() == "true"
    valor = int(float(check.get("paid_amount") or check.get("amount") or 0))
    if not pago:
        return "nao_pago"
    if valor < pedido["centavos"]:  # juros do parcelamento só aumentam o valor
        log.error("Pedido %s pago com valor menor (%s < %s)", pedido_id, valor, pedido["centavos"])
        return "nao_pago"

    # Marca o pedido como pago antes de creditar: se o aviso e o retorno do
    # cliente chegarem juntos, só um deles consegue esta gravação.
    pedido.update(status="pago", transacao=transacao, slug=slug, pago_em=time.time(), valor_pago=valor)
    try:
        gravar_json_versionado(caminho, pedido, geracao)
    except PreconditionFailed:
        return "ja_pago"

    def creditar(usuarios):
        conta = usuarios.get(pedido["usuario"])
        if conta is None:
            log.error("Pedido %s pago para conta inexistente %s", pedido_id, pedido["usuario"])
            return False
        conta["pago_ate"] = max(time.time(), conta.get("pago_ate", 0)) + pedido["dias"] * DIA
        conta["plano"] = pedido["plano"]
        if conta.get("expira_em") or conta.get("teste_acabou_em"):   # assinou durante/depois do teste
            conta.pop("expira_em", None)
            conta["assinante"] = True
        return conta

    try:
        conta = alterar_usuarios(creditar)
    except Exception:
        log.exception("PAGAMENTO SEM CRÉDITO: pedido %s de %s — liberar à mão no /admin", pedido_id, pedido["usuario"])
        raise
    if conta:
        somar_metricas(conta.get("origem"), pagamentos=1, receita_centavos=pedido["centavos"])
        enviar_boas_vindas(email_da_conta(pedido["usuario"], conta) or "", dict(conta), pedido["usuario"])
    if conta and has_request_context() and session.get("usuario") == pedido["usuario"]:
        atualizar_sessao_assinatura(conta)
        session["expira_em"] = None
    return "pago"


# --------------------------------------------------------------------------- #
# Origem dos visitantes (links com ?origem=professores, ?origem=pais...): quantos
# visitaram, testaram o mini-quiz, criaram conta e pagaram, por origem. Fica em
# metricas/origens.json no Firebase e aparece no /admin.
# --------------------------------------------------------------------------- #
ARQUIVO_METRICAS = "metricas/origens.json"
COOKIE_ORIGEM = "qz_origem"
ORIGEM_VALIDA = re.compile(r"[a-z0-9_-]{1,40}")
MAX_ORIGENS = 200                                  # evita que links inventados encham o arquivo
# Pré-visualizações de link (Facebook, WhatsApp...) e robôs não contam como visita
ROBOS = re.compile(r"bot|crawl|spider|facebookexternalhit|facebot|preview|whatsapp|telegram|slack|discord", re.I)


def origem_valida(valor):
    valor = (valor or "").strip().lower()
    return valor if ORIGEM_VALIDA.fullmatch(valor) else None


def eh_robo():
    return bool(ROBOS.search(request.headers.get("User-Agent", "")))


def origem_atual():
    """O ?origem= desta visita, ou o de uma visita anterior (cookie); sem nenhum,
    'site' quando veio da landing principal e 'direto' nos outros casos."""
    origem = origem_valida(request.args.get("origem")) or origem_valida(request.cookies.get(COOKIE_ORIGEM))
    if origem:
        return origem
    referer = request.headers.get("Referer", "")
    return "site" if re.match(r"https?://(www\.)?quizzify\.com\.br", referer) else "direto"


def _somar_metricas(origem, valores):
    from google.api_core.exceptions import PreconditionFailed
    for _ in range(6):
        try:
            dados, geracao = ler_json_versionado(ARQUIVO_METRICAS)
            dados = dados or {}
            if origem not in dados and len(dados) >= MAX_ORIGENS:
                origem = "outros"
            linha = dados.setdefault(origem, {})
            for campo, n in valores.items():
                linha[campo] = linha.get(campo, 0) + n
            gravar_json_versionado(ARQUIVO_METRICAS, dados, geracao)
            return
        except PreconditionFailed:
            time.sleep(0.2)                        # outra gravação ao mesmo tempo: tenta de novo
        except Exception:
            log.exception("Falha ao somar métricas de %s", origem)
            return


def somar_metricas(origem, **valores):
    """Soma em segundo plano, para não atrasar a página."""
    threading.Thread(target=_somar_metricas, args=(origem or "direto", valores), daemon=True).start()


def ler_metricas():
    try:
        dados, _ = ler_json_versionado(ARQUIVO_METRICAS)
    except Exception:
        log.exception("Falha ao ler métricas")
        return None
    linhas = [{"origem": o, **{c: v.get(c, 0) for c in ("visitantes", "testes", "contas", "pagamentos", "receita_centavos")}}
              for o, v in (dados or {}).items()]
    return sorted(linhas, key=lambda l: (-l["pagamentos"], -l["contas"], -l["visitantes"]))


@app.before_request
def registrar_origem():
    if request.method != "GET" or request.endpoint in (None, "static", "health", "service_worker", "manifesto",
                                                       "demo", "demo_video", "tarefa_lembretes", "divulgacao_js") or eh_robo():
        return
    pedida = origem_valida(request.args.get("origem"))
    guardada = origem_valida(request.cookies.get(COOKIE_ORIGEM))
    if pedida and pedida != guardada:
        g.nova_origem = pedida
    elif not pedida and not guardada and request.endpoint == "assine":
        g.nova_origem = origem_atual()           # chegou à /assine sem link marcado
    if g.get("nova_origem"):
        somar_metricas(g.nova_origem, visitantes=1)


@app.after_request
def gravar_cookie_origem(resposta):
    if g.get("nova_origem"):
        resposta.set_cookie(COOKIE_ORIGEM, g.nova_origem, max_age=30 * DIA, samesite="Lax", httponly=True,
                            secure=bool(os.environ.get("RENDER")))
        resposta.headers["Cache-Control"] = "private, max-age=600"
    return resposta


@app.post("/m/teste")
def metrica_teste():
    """A /assine avisa quando a pessoa abre o mini-quiz."""
    if not eh_robo():
        somar_metricas(origem_atual(), testes=1)
    return "", 204


# --------------------------------------------------------------------------- #
# Rotas
# --------------------------------------------------------------------------- #
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html")

    ip = ip_cliente()
    if bloqueado(ip):
        flash("Muitas tentativas. Aguarde alguns minutos e tente novamente.", "erro")
        return render_template("login.html"), 429

    nome = request.form.get("usuario", "").strip().lower()
    senha = request.form.get("senha", "")
    usuarios = carregar_usuarios()
    conta = usuarios.get(nome)
    if not conta or not check_password_hash(conta["senha"], senha):
        registrar_tentativa(_tentativas, ip)
        flash("Utilizador ou senha incorretos.", "erro")
        return render_template("login.html", usuario=nome), 401
    with _trava_contadores:
        _tentativas.pop(ip, None)
    if expirou(conta):
        try:
            conta = converter_teste_expirado(nome) or conta
        except Exception:
            log.exception("Falha ao converter teste de %s", nome)
            flash("Não foi possível entrar agora. Tente novamente.", "erro")
            return render_template("login.html", usuario=nome), 503
        entrar_na_sessao(nome, conta)
        flash("O seu teste acabou. Escolha um plano para continuar — os seus quizzes estão guardados. 💾", "erro")
        return redirect(url_for("assinatura", plano="anual"))
    entrar_na_sessao(nome, conta)
    proximo = request.args.get("proximo", "")
    return redirect(proximo if proximo.startswith("/") and not proximo.startswith("//") else url_for("index"))


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/trocar-senha", methods=["GET", "POST"])
@login_obrigatorio
def trocar_senha():
    if request.method == "POST":
        try:
            usuarios = carregar_usuarios(estrito=True)
        except FalhaUsuarios:
            flash("Serviço temporariamente indisponível. Tente novamente.", "erro")
            return render_template("trocar_senha.html", obrigatorio=session.get("trocar_senha")), 503
        if not conta_ainda_existe(usuarios):
            return redirect(url_for("login"))
        conta = usuarios[session["usuario"]]
        atual, nova = request.form.get("atual", ""), request.form.get("nova", "")
        if not conta or not check_password_hash(conta["senha"], atual):
            flash("Senha atual incorreta.", "erro")
        elif not senha_valida(nova):
            flash("A nova senha deve ter pelo menos 8 caracteres.", "erro")
        elif nova != request.form.get("confirmar", ""):
            flash("A confirmação não coincide com a nova senha.", "erro")
        elif nova == atual:
            flash("A nova senha deve ser diferente da atual.", "erro")
        else:
            conta.update(senha=generate_password_hash(nova), trocar_senha=False)
            try:
                salvar_usuarios(usuarios)
            except Exception:
                log.exception("Falha ao guardar utilizadores")
                flash("Não foi possível guardar a senha. Tente novamente.", "erro")
            else:
                session["trocar_senha"] = False
                flash("Senha alterada com sucesso.", "ok")
                return redirect(url_for("index"))
    return render_template("trocar_senha.html", obrigatorio=session.get("trocar_senha"))


@app.route("/admin", methods=["GET", "POST"])
@admin_obrigatorio
def admin():
    try:
        usuarios = carregar_usuarios(estrito=True)
    except FalhaUsuarios:
        flash("Não foi possível ler os utilizadores do Firebase. Tente novamente.", "erro")
        return render_template("admin.html", usuarios=[]), 503
    if not conta_ainda_existe(usuarios) or not usuarios[session["usuario"]].get("admin"):
        session.clear()
        return redirect(url_for("login"))
    if request.method == "POST":
        acao = request.form.get("acao")
        nome = request.form.get("usuario", "").strip().lower()
        senha = request.form.get("senha", "")
        if acao == "criar":
            if not (re.fullmatch(r"[a-z0-9._-]{3,30}", nome) or (EMAIL_VALIDO.fullmatch(nome) and len(nome) <= 120)):
                flash("Use o e-mail da pessoa (recomendado) ou um nome de 3 a 30 letras minúsculas, números, ponto, hífen ou _.", "erro")
            elif nome in usuarios:
                flash(f"O utilizador '{nome}' já existe.", "erro")
            elif not senha_valida(senha):
                flash("A senha deve ter pelo menos 8 caracteres.", "erro")
            elif (horas := DURACOES_TESTE.get(request.form.get("duracao"))):
                # Conta de teste: sem troca de senha obrigatória, nunca admin, expira sozinha
                expira = time.time() + horas * 3600
                usuarios[nome] = {"senha": generate_password_hash(senha), "admin": False,
                                  "trocar_senha": False, "expira_em": expira, "teste_horas": horas}
                flash(f"Acesso de teste '{nome}' criado. Expira em {data_hora(expira)}.", "ok")
            else:
                usuarios[nome] = {"senha": generate_password_hash(senha),
                                  "admin": request.form.get("admin") == "1", "trocar_senha": True}
                flash(f"Utilizador '{nome}' criado. Ele terá de trocar a senha no primeiro acesso.", "ok")
        elif acao == "redefinir" and nome in usuarios:
            if not senha_valida(senha):
                flash("A senha deve ter pelo menos 8 caracteres.", "erro")
            else:
                usuarios[nome].update(senha=generate_password_hash(senha),
                                      trocar_senha=not usuarios[nome].get("expira_em"))
                flash(f"Senha de '{nome}' redefinida.", "ok")
        elif acao == "renovar" and (usuarios.get(nome, {}).get("expira_em") or (
                usuarios.get(nome, {}).get("teste_acabou_em") and not usuarios[nome].get("pago_ate"))):
            conta = usuarios[nome]
            if not conta.get("expira_em"):        # teste já tinha acabado: volta a ser teste
                conta.pop("assinante", None)
                conta.pop("teste_acabou_em", None)
                conta["teste_horas"] = 24
            expira = max(time.time(), conta.get("expira_em") or 0) + 24 * 3600
            conta["expira_em"] = expira
            flash(f"Acesso de '{nome}' prolongado até {data_hora(expira)}.", "ok")
        elif acao == "liberar30" and usuarios.get(nome, {}).get("assinante"):
            conta = usuarios[nome]
            conta["pago_ate"] = max(time.time(), conta.get("pago_ate", 0)) + 30 * DIA
            flash(f"Assinatura de '{nome}' liberada até {data_hora(conta['pago_ate'])}.", "ok")
        elif acao == "divulgacao":
            novo = {"ligada": request.form.get("ligada") == "1",
                    "texto": request.form.get("texto", "").strip()[:160] or DIVULGACAO_PADRAO["texto"],
                    "destaque": request.form.get("destaque", "").strip()[:160],
                    "botao": request.form.get("botao", "").strip()[:60] or DIVULGACAO_PADRAO["botao"]}
            try:
                gravar_json_versionado(ARQUIVO_DIVULGACAO, novo)
                _divulgacao_cache.update(dados=novo, quando=time.time())
                flash(f"Faixa de divulgação {'ligada' if novo['ligada'] else 'desligada'}. Vale para todos os quizzes em até 5 minutos.", "ok")
            except Exception:
                log.exception("Falha ao salvar divulgação")
                flash("Não foi possível salvar a faixa agora. Tente novamente.", "erro")
            return redirect(url_for("admin"))
        elif acao == "email" and nome in usuarios:
            email = request.form.get("email", "").strip().lower()
            if email and not EMAIL_VALIDO.fullmatch(email):
                flash("E-mail inválido.", "erro")
            else:
                if email:
                    usuarios[nome]["email"] = email
                else:
                    usuarios[nome].pop("email", None)
                flash(f"E-mail de '{nome}' {'salvo' if email else 'removido'}.", "ok")
        elif acao == "remover" and nome in usuarios:
            if nome == session["usuario"]:
                flash("Não pode remover a sua própria conta.", "erro")
            else:
                del usuarios[nome]
                flash(f"Utilizador '{nome}' removido.", "ok")
        else:
            abort(400)

        try:
            salvar_usuarios(usuarios)
        except Exception:
            log.exception("Falha ao guardar utilizadores")
            flash("Não foi possível guardar as alterações no Firebase.", "erro")
        return redirect(url_for("admin"))

    return render_template("admin.html", usuarios=sorted(usuarios.items()), agora=time.time(),
                           config_limite_teste=LIMITE_TESTE, metricas=ler_metricas(),
                           divulgacao=ler_divulgacao(fresco=True))


EMAIL_VALIDO = re.compile(r"[^@\s]+@[^@\s]+\.[a-z]{2,}")
_cadastros = {}


def plano_escolhido():
    """Plano vindo da landing page (?plano=mensal|anual), ou None."""
    plano = (request.values.get("plano") or "").strip().lower()
    return plano if plano in PLANOS else None


@app.route("/cadastro", methods=["GET", "POST"])
def cadastro():
    plano_id = plano_escolhido()
    if session.get("usuario"):
        return redirect(url_for("assinatura", plano=plano_id) if plano_id else url_for("index"))
    if request.method == "GET":
        return render_template("cadastro.html", plano_id=plano_id, planos=PLANOS)

    email = request.form.get("email", "").strip().lower()
    senha = request.form.get("senha", "")
    ip, agora = ip_cliente(), time.time()
    with _trava_contadores:
        _cadastros[ip] = [t for t in _cadastros.get(ip, []) if agora - t < 3600]
        excedeu = len(_cadastros[ip]) >= 5
    if excedeu:
        flash("Muitos cadastros a partir desta ligação. Tente mais tarde.", "erro")
    elif not EMAIL_VALIDO.fullmatch(email) or len(email) > 120:
        flash("Informe um e-mail válido.", "erro")
    elif not senha_valida(senha):
        flash("A senha deve ter pelo menos 8 caracteres.", "erro")
    elif senha != request.form.get("confirmar", ""):
        flash("A confirmação não coincide com a senha.", "erro")
    else:
        nova = {"senha": generate_password_hash(senha), "admin": False, "trocar_senha": False,
                "assinante": True, "pago_ate": 0, "criado_em": agora, "origem": origem_atual()}

        def criar(usuarios):
            if email in usuarios:
                return False
            usuarios[email] = nova
            return nova

        try:
            criada = alterar_usuarios(criar)
        except Exception:
            log.exception("Falha ao criar conta")
            flash("Não foi possível criar a conta agora. Tente novamente.", "erro")
            return render_template("cadastro.html", email=email, plano_id=plano_id, planos=PLANOS), 503
        if criada is not False:
            somar_metricas(nova["origem"], contas=1)
        if criada is False:
            flash("Já existe uma conta com este e-mail. Entre com a sua senha.", "erro")
            return redirect(url_for("login", proximo=url_for("assinatura", plano=plano_id) if plano_id else None))
        registrar_tentativa(_cadastros, ip, agora)
        session.clear()
        session.permanent = True
        session.update(usuario=email, admin=False, trocar_senha=False, expira_em=None)
        atualizar_sessao_assinatura(nova)
        # Veio da landing com um plano escolhido: vai direto para o pagamento dele.
        if plano_id and INFINITEPAY_HANDLE:
            try:
                return redirect(iniciar_checkout(plano_id))
            except Exception:
                log.exception("Falha ao criar checkout após cadastro")
                flash("Conta criada! Não foi possível abrir o pagamento agora; tente novamente abaixo.", "erro")
        return redirect(url_for("assinatura", plano=plano_id))
    return render_template("cadastro.html", email=email, plano_id=plano_id, planos=PLANOS), 400


@app.get("/assinatura")
@login_obrigatorio
def assinatura():
    conta = carregar_usuarios().get(session["usuario"])
    if conta:
        atualizar_sessao_assinatura(conta)
    return render_template("assinatura.html", conta=conta or {}, planos=PLANOS, agora=time.time(),
                           pagamentos_ativos=bool(INFINITEPAY_HANDLE),
                           total_quizzes=len(listar_historico(session["usuario"])),
                           destaque=plano_escolhido() or "anual")


def iniciar_checkout(plano_id):
    """Cria o pedido e devolve o link de pagamento da InfinitePay."""
    plano = PLANOS[plano_id]
    pedido_id = uuid.uuid4().hex
    pedido = {"usuario": session["usuario"], "plano": plano_id, "centavos": plano["centavos"],
              "dias": plano["dias"], "status": "pendente", "criado_em": time.time()}
    gravar_json_versionado(f"pedidos/{pedido_id}.json", pedido, 0)
    return criar_checkout(pedido_id, plano, session["usuario"])


@app.post("/assinatura/pagar")
@login_obrigatorio
def assinatura_pagar():
    plano_id = request.form.get("plano")
    if plano_id not in PLANOS or not INFINITEPAY_HANDLE or not (session.get("assinante") or session.get("expira_em")):
        abort(400)
    try:
        link = iniciar_checkout(plano_id)
    except Exception:
        log.exception("Falha ao criar checkout")
        flash("Não foi possível abrir o pagamento agora. Tente novamente em instantes.", "erro")
        return redirect(url_for("assinatura"))
    return redirect(link)


@app.get("/assinatura/retorno")
@login_obrigatorio
def assinatura_retorno():
    """A InfinitePay devolve o cliente aqui depois de pagar. Confirmamos já,
    sem esperar pelo aviso, para o acesso abrir na hora."""
    pedido_id = request.args.get("order_nsu", "")
    transacao = request.args.get("transaction_nsu", "")
    slug = request.args.get("slug") or request.args.get("invoice_slug", "")
    if pedido_id and transacao and slug:
        try:
            resultado = confirmar_pagamento(pedido_id, transacao, slug)
        except Exception:
            log.exception("Falha ao confirmar retorno do pedido %s", pedido_id)
            resultado = "erro"
        if resultado in ("pago", "ja_pago"):
            flash("Pagamento confirmado! O seu acesso está liberado. 🎉", "ok")
            return redirect(url_for("assinatura"))
    flash("Ainda não recebemos a confirmação do pagamento. Se já pagou, atualize esta página "
          "em alguns minutos.", "erro")
    return redirect(url_for("assinatura"))


@app.post("/webhooks/infinitepay")
def webhook_infinitepay():
    if not hmac.compare_digest(request.args.get("token", ""), token_webhook()):
        abort(403)
    dados = request.get_json(silent=True) or {}
    pedido_id = str(dados.get("order_nsu") or "")
    transacao = str(dados.get("transaction_nsu") or "")
    slug = str(dados.get("invoice_slug") or dados.get("slug") or "")
    if not (pedido_id and transacao and slug):
        return jsonify(ok=False, erro="dados incompletos"), 200  # nada a repetir
    try:
        resultado = confirmar_pagamento(pedido_id, transacao, slug)
    except Exception:
        log.exception("Falha ao processar aviso do pedido %s", pedido_id)
        return jsonify(ok=False), 400  # 400 faz a InfinitePay reenviar
    log.info("Aviso InfinitePay pedido=%s resultado=%s", pedido_id, resultado)
    return jsonify(ok=True, resultado=resultado), 200


# --------------------------------------------------------------------------- #
# Lembretes de renovação por e-mail (Brevo)
# Um cron-job chama /tarefas/lembretes?token=... uma vez por dia. Cada assinante
# recebe no máximo um e-mail "vence em breve" (últimos 5 dias) e um "venceu"
# (até 3 dias depois) por período pago; renovar zera o controle.
# --------------------------------------------------------------------------- #
BREVO_API = "https://api.brevo.com/v3/smtp/email"
EMAIL_REMETENTE = os.environ.get("EMAIL_REMETENTE", "").strip()
EMAIL_REMETENTE_NOME = os.environ.get("EMAIL_REMETENTE_NOME", "Quizzify").strip()
JANELA_VENCIDA_DIAS = 3


EMAIL_SUPORTE = os.environ.get("EMAIL_SUPORTE", "contato@quizzify.com.br").strip()
# Links dos e-mails: sempre o domínio próprio (o aviso de pagamento chega pelo endereço do Render)
URL_APP = (os.environ.get("URL_PUBLICA", "").strip() or "https://app.quizzify.com.br").rstrip("/")


def enviar_email(para, assunto, html_corpo, texto_corpo, tag="lembrete-renovacao"):
    # HTML + versão em texto simples: os filtros de spam penalizam e-mails só em HTML.
    resposta = requests.post(BREVO_API, timeout=15, headers={
        "api-key": obter_segredo("BREVO_API_KEY"), "accept": "application/json",
        "content-type": "application/json"}, json={
        "sender": {"name": EMAIL_REMETENTE_NOME, "email": EMAIL_REMETENTE},
        "replyTo": {"email": EMAIL_SUPORTE, "name": "Suporte Quizzify"},
        "to": [{"email": para}], "subject": assunto, "htmlContent": html_corpo,
        "textContent": texto_corpo, "tags": [tag]})
    if resposta.status_code not in (200, 201, 202):
        raise RuntimeError(f"Brevo devolveu {resposta.status_code}: {resposta.text[:300]}")


def email_lembrete(etapa, conta, total_quizzes):
    """Devolve (assunto, html, texto simples) do lembrete."""
    plano = PLANOS.get(conta.get("plano"), {}).get("nome", "")
    quando = data_hora(conta["pago_ate"])
    link = url_publica("assinatura")
    if etapa == "antes":
        dias = max(1, round((conta["pago_ate"] - time.time()) / DIA))
        assunto = f"Sua assinatura do Quizzify vence em {dias} dia{'s' if dias > 1 else ''}"
        titulo = "⏳ Sua assinatura está perto de vencer"
        texto = (f"Sua assinatura{' do plano ' + plano if plano else ''} vence em <b>{quando}</b>. "
                 "Renove agora para continuar criando quizzes sem interrupção — "
                 "os dias que ainda faltam são somados ao novo período.")
    else:
        assunto = "Sua assinatura do Quizzify venceu"
        titulo = "⛔ Sua assinatura venceu"
        salvos = (f" Seus <b>{total_quizzes} quiz{'zes' if total_quizzes > 1 else ''}</b> continuam guardados"
                  " e voltam a ficar disponíveis assim que renovar.") if total_quizzes else ""
        texto = f"Sua assinatura venceu em <b>{quando}</b> e o acesso foi pausado.{salvos}"
    corpo = f"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{assunto}</title></head>
<body style="margin:0;padding:0;background:#ffffff">
<div style="font-family:Arial,sans-serif;max-width:480px;margin:auto;padding:24px;color:#1e293b">
<p style="font-family:Arial,sans-serif;font-size:30px;font-weight:800;letter-spacing:-1px;color:#24133E;margin:0 0 16px">quizzify<span style="color:#5920AB">.</span></p>
<h2 style="font-size:18px;margin:0 0 12px">{titulo}</h2>
<p style="font-size:15px;line-height:1.5">{texto}</p>
<p style="text-align:center;margin:28px 0"><a href="{link}" style="background:#5920AB;color:#fff;text-decoration:none;
font-weight:bold;padding:14px 28px;border-radius:12px;display:inline-block">Renovar assinatura</a></p>
<p style="font-size:13px;color:#64748b">Pagamento por Pix ou cartão, com liberação imediata.<br>
Se já renovou, pode ignorar este e-mail.<br>
Dúvidas? Responda este e-mail ou escreva para <a href="mailto:{EMAIL_SUPORTE}" style="color:#64748b">{EMAIL_SUPORTE}</a>.</p></div>
</body></html>"""
    texto_simples = (f"Quizzify\n\n{titulo}\n\n{re.sub(r'<[^>]+>', '', texto)}\n\n"
                     f"Renovar assinatura: {link}\n\n"
                     "Pagamento por Pix ou cartão, com liberação imediata.\n"
                     "Se já renovou, pode ignorar este e-mail.\n"
                     f"Dúvidas? Escreva para {EMAIL_SUPORTE}\n")
    return assunto, corpo, texto_simples


def montar_email(assunto, titulo, miolo_html, botao, link, miolo_texto, depois_html=""):
    """Layout comum dos e-mails (HTML simples, funciona em todos os leitores)."""
    corpo = f"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{assunto}</title></head>
<body style="margin:0;padding:0;background:#ffffff">
<div style="font-family:Arial,sans-serif;max-width:480px;margin:auto;padding:24px;color:#1e293b">
<p style="font-family:Arial,sans-serif;font-size:30px;font-weight:800;letter-spacing:-1px;color:#24133E;margin:0 0 16px">quizzify<span style="color:#5920AB">.</span></p>
<h2 style="font-size:18px;margin:0 0 12px">{titulo}</h2>
{miolo_html}
<p style="text-align:center;margin:28px 0"><a href="{link}" style="background:#5920AB;color:#fff;text-decoration:none;
font-weight:bold;padding:14px 28px;border-radius:12px;display:inline-block">{botao}</a></p>
{depois_html}
<p style="font-size:13px;color:#64748b">Dúvidas ou dificuldade para acessar? Responda este e-mail ou escreva para
<a href="mailto:{EMAIL_SUPORTE}" style="color:#64748b">{EMAIL_SUPORTE}</a>.</p></div>
</body></html>"""
    texto = f"Quizzify\n\n{titulo}\n\n{miolo_texto}\n\n{botao}: {link}\n\nDúvidas? Escreva para {EMAIL_SUPORTE}\n"
    return assunto, corpo, texto


def email_boas_vindas(email, conta):
    plano = PLANOS.get(conta.get("plano"), {}).get("nome", "")
    validade = datetime.fromtimestamp(conta["pago_ate"], FUSO_BRASIL).strftime("%d/%m/%Y")
    p = 'style="font-size:15px;line-height:1.5"'
    miolo = f"""<p {p}>Pagamento confirmado! Seu acesso{' ao plano <b>' + plano + '</b>' if plano else ''}
está liberado até <b>{validade}</b>.</p>
<p {p}><b>Para entrar:</b> acesse <a href="{URL_APP}">app.quizzify.com.br</a> com
<b>{email}</b> e a sua senha.</p>"""
    depois = f"""<p {p}><b>📲 Quer o Quizzify na tela do celular?</b> (opcional — funciona igual no navegador)<br>
<b>Android:</b> abra o link no <b>Chrome</b>, toque em <b>⋮</b> e em <b>Instalar app</b>.<br>
<b>iPhone:</b> abra no <b>Safari</b>, toque em <b>Compartilhar</b> (□↑) e em <b>Adicionar à Tela de Início</b>.</p>
<p style="font-size:13px;color:#64748b">Esqueceu a senha? Na tela de entrar, toque em
<a href="{URL_APP}/esqueci-senha" style="color:#64748b">Esqueci minha senha</a>.</p>"""
    texto = (f"Pagamento confirmado! Seu acesso{' ao plano ' + plano if plano else ''} está liberado até {validade}.\n\n"
             f"Para entrar: acesse {URL_APP} com {email} e a sua senha.\n\n"
             "Quer o Quizzify na tela do celular? (opcional)\n"
             "Android: abra o link no Chrome, toque em ⋮ e em Instalar app.\n"
             "iPhone: abra no Safari, toque em Compartilhar e em Adicionar à Tela de Início.\n\n"
             f"Esqueceu a senha? {URL_APP}/esqueci-senha")
    return montar_email("Seu acesso ao Quizzify está liberado 🎉", "🎉 Bem-vindo(a) ao Quizzify!",
                        miolo, "Acessar o Quizzify", URL_APP + "/", texto, depois)


def enviar_boas_vindas(email, conta, nome=None):
    """Roda em segundo plano depois do pagamento; uma falha aqui não afeta o crédito."""
    if "@" not in email or not os.environ.get("BREVO_API_KEY") or not EMAIL_REMETENTE:
        return
    def tarefa():
        try:
            enviar_email(email, *email_boas_vindas(nome or email, conta), tag="boas-vindas")
        except Exception:
            log.exception("Falha ao enviar boas-vindas para %s", email)
    threading.Thread(target=tarefa, daemon=True).start()


# ---------- Esqueci minha senha ----------
VALIDADE_LINK_SENHA = 3600
_pedidos_senha = {}


def _assinador_senha():
    return URLSafeTimedSerializer(app.secret_key, salt="nova-senha")


def email_nova_senha(link):
    miolo = ('<p style="font-size:15px;line-height:1.5">Recebemos um pedido para criar uma nova senha. '
             'Toque no botão abaixo — o link vale por <b>1 hora</b> e só funciona uma vez.</p>'
             '<p style="font-size:13px;color:#64748b">Não pediu? Pode ignorar este e-mail: sua senha continua a mesma.</p>')
    texto = ("Recebemos um pedido para criar uma nova senha. O link vale por 1 hora e só funciona uma vez.\n"
             "Não pediu? Pode ignorar este e-mail.")
    return montar_email("Criar nova senha do Quizzify", "🔑 Criar nova senha", miolo, "Criar nova senha", link, texto)


@app.route("/esqueci-senha", methods=["GET", "POST"])
def esqueci_senha():
    if request.method == "GET":
        return render_template("esqueci_senha.html")
    email = request.form.get("email", "").strip().lower()
    ip, agora = ip_cliente(), time.time()
    with _trava_contadores:
        _pedidos_senha[ip] = [t for t in _pedidos_senha.get(ip, []) if agora - t < 3600]
        excedeu = len(_pedidos_senha[ip]) >= 5
    if excedeu:
        flash("Muitos pedidos. Aguarde um pouco e tente novamente.", "erro")
        return render_template("esqueci_senha.html", email=email), 429
    registrar_tentativa(_pedidos_senha, ip, agora)
    nome, conta = None, None
    if EMAIL_VALIDO.fullmatch(email):
        usuarios = carregar_usuarios()
        nome = email if email in usuarios else next(
            (n for n, c in usuarios.items() if (c.get("email") or "").lower() == email), None)
        conta = usuarios.get(nome) if nome else None
    if conta and os.environ.get("BREVO_API_KEY") and EMAIL_REMETENTE:
        token = _assinador_senha().dumps({"u": nome, "h": conta["senha"][-16:]})
        link = f"{URL_APP}/nova-senha/{token}"
        def tarefa():
            try:
                enviar_email(email, *email_nova_senha(link), tag="nova-senha")
            except Exception:
                log.exception("Falha ao enviar link de nova senha para %s", email)
        threading.Thread(target=tarefa, daemon=True).start()
    # Mesma resposta sempre: não revela quais e-mails têm conta.
    return render_template("esqueci_senha.html", enviado=email)


@app.route("/nova-senha/<token>", methods=["GET", "POST"])
def nova_senha(token):
    try:
        dados = _assinador_senha().loads(token, max_age=VALIDADE_LINK_SENHA)
    except BadSignature:
        dados = {}
    nome = dados.get("u", "")
    conta = carregar_usuarios().get(nome) if nome else None
    if not conta or conta["senha"][-16:] != dados.get("h"):
        flash("Este link expirou ou já foi usado. Peça um novo abaixo.", "erro")
        return redirect(url_for("esqueci_senha"))
    if request.method == "GET":
        return render_template("nova_senha.html", email=nome)


    nova = request.form.get("nova", "")
    if not senha_valida(nova):
        flash("A nova senha deve ter pelo menos 8 caracteres.", "erro")
    elif nova != request.form.get("confirmar", ""):
        flash("A confirmação não coincide com a nova senha.", "erro")
    else:
        def trocar(usuarios):
            c = usuarios.get(nome)
            if not c or c["senha"][-16:] != dados.get("h"):
                return False
            c.update(senha=generate_password_hash(nova), trocar_senha=False)
            return c
        try:
            atualizada = alterar_usuarios(trocar)
        except Exception:
            log.exception("Falha ao gravar nova senha")
            flash("Não foi possível salvar agora. Tente novamente.", "erro")
            return render_template("nova_senha.html", email=nome), 503
        if atualizada is False:
            flash("Este link já foi usado. Peça um novo abaixo.", "erro")
            return redirect(url_for("esqueci_senha"))
        entrar_na_sessao(nome, atualizada)
        flash("Senha alterada! Você já está dentro. 🎉", "ok")
        return redirect(url_for("index"))
    return render_template("nova_senha.html", email=nome), 400


AVISO_TESTE_DIAS = 2


def etapa_teste(conta, agora):
    """Aviso de fim do teste: "teste_antes" (até 2 dias antes) ou "teste_fim" (até 3 dias depois)."""
    expira = conta.get("expira_em") or 0
    if not expira or conta.get("admin"):
        return None
    restante = expira - agora
    if 0 < restante <= AVISO_TESTE_DIAS * DIA:
        if (conta.get("teste_horas") or 999) <= 24:
            return None                     # teste de 24h: só o aviso no app e o e-mail do fim
        etapa = "teste_antes"
    elif -JANELA_VENCIDA_DIAS * DIA < restante <= 0:
        etapa = "teste_fim"
    else:
        return None
    registro = conta.get("lembretes_teste") or {}
    if registro.get("ref") == expira and etapa in registro.get("enviados", []):
        return None
    return etapa


def email_teste(etapa, nome, conta, total_quizzes):
    mensal, anual = PLANOS["mensal"]["centavos"], PLANOS["anual"]["centavos"]
    reais = lambda c: f"R$ {c / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    precos = (f"<b>Mensal:</b> {reais(mensal)}/mês &nbsp;·&nbsp; <b>Anual:</b> {reais(anual)}/ano "
              f"(sai por {reais(anual / 12)}/mês)")
    precos_txt = re.sub(r"<[^>]+>|&nbsp;", "", precos)
    p = 'style="font-size:15px;line-height:1.5"'
    entrar = (f'<p style="font-size:13px;color:#64748b">Para entrar, use <b>{nome}</b> e a sua senha.</p>')
    guardados = (f" Seus <b>{total_quizzes} quiz{'zes' if total_quizzes != 1 else ''}</b> continuam guardados."
                 if total_quizzes else " Tudo o que você criou continua guardado.")
    if etapa == "teste_antes":
        dias = int(max(0, conta["expira_em"] - time.time()) // DIA) + 1
        assunto = f"Seu teste do Quizzify acaba em {dias} dia{'s' if dias > 1 else ''}"
        titulo = "⏳ Seu teste está acabando"
        texto = (f"Seu acesso de teste termina em <b>{data_hora(conta['expira_em'])}</b>. "
                 f"Para continuar criando quizzes e flashcards, escolha um plano.{guardados}")
        botao = "Escolher meu plano"
    else:
        assunto = "Seu teste do Quizzify acabou — seus quizzes estão guardados"
        titulo = "Seu teste acabou 💾"
        texto = f"Esperamos que tenha gostado!{guardados} Assine para voltar a criar quizzes na hora."
        botao = "Assinar e continuar"
    miolo = f"<p {p}>{texto}</p><p {p}>{precos}<br>Pix ou cartão, liberação imediata.</p>"
    miolo_txt = f"{re.sub(r'<[^>]+>', '', texto)}\n\n{precos_txt}\nPix ou cartão, liberação imediata.\n\nPara entrar, use {nome} e a sua senha."
    return montar_email(assunto, titulo, miolo, botao, f"{URL_APP}/assinatura?plano=anual", miolo_txt, entrar)


def etapa_lembrete(conta, agora):
    """Qual lembrete esta conta deve receber agora (ou None)."""
    pago_ate = conta.get("pago_ate") or 0
    if not conta.get("assinante") or not pago_ate:
        return None
    restante = pago_ate - agora
    if 0 < restante <= AVISO_VENCIMENTO_DIAS * DIA:
        etapa = "antes"
    elif -JANELA_VENCIDA_DIAS * DIA < restante <= 0:
        etapa = "vencida"
    else:
        return None
    registro = conta.get("lembretes") or {}
    if registro.get("ref") == pago_ate and etapa in registro.get("enviados", []):
        return None
    return etapa


@app.route("/tarefas/lembretes", methods=["GET", "POST"])
def tarefa_lembretes():
    token = os.environ.get("LEMBRETES_TOKEN", "")
    if not token:
        return jsonify(erro="LEMBRETES_TOKEN não configurado"), 503
    if not hmac.compare_digest(request.args.get("token", ""), token):
        abort(403)
    if not os.environ.get("BREVO_API_KEY") or not EMAIL_REMETENTE:
        return jsonify(erro="BREVO_API_KEY ou EMAIL_REMETENTE não configurados"), 503

    # ?teste=email@exemplo.com envia um lembrete de exemplo, sem mexer em contas
    if (teste := request.args.get("teste", "").strip()):
        exemplo = {"pago_ate": time.time() + 3 * DIA, "plano": "mensal"}
        try:
            if request.args.get("tipo") in ("teste_antes", "teste_fim"):
                exemplo = {"expira_em": time.time() + (2 * DIA if request.args["tipo"] == "teste_antes" else -3600)}
                enviar_email(teste, *email_teste(request.args["tipo"], teste, exemplo, 3), tag="fim-do-teste")
            elif request.args.get("tipo") == "boas-vindas":   # &tipo=boas-vindas: e-mail pós-pagamento
                exemplo = {"pago_ate": time.time() + 365 * DIA, "plano": "anual"}
                enviar_email(teste, *email_boas_vindas(teste, exemplo), tag="boas-vindas")
            else:
                enviar_email(teste, *email_lembrete("antes", exemplo, 0))
        except Exception as erro:
            log.exception("Falha no e-mail de teste")
            return jsonify(ok=False, erro=str(erro)), 502
        return jsonify(ok=True, enviado_para=teste)

    try:
        usuarios = carregar_usuarios(estrito=True)
    except FalhaUsuarios:
        return jsonify(erro="Falha ao ler utilizadores"), 503

    agora, enviados, enviados_teste, falhas = time.time(), [], [], 0
    for nome, conta in usuarios.items():
        destino = email_da_conta(nome, conta)
        if not destino:
            continue
        if (etapa := etapa_teste(conta, agora)):
            try:
                enviar_email(destino, *email_teste(etapa, nome, conta, len(listar_historico(nome))), tag="fim-do-teste")
                enviados_teste.append((nome, etapa, conta["expira_em"]))
            except Exception:
                log.exception("Falha ao enviar aviso de teste '%s' para %s", etapa, nome)
                falhas += 1
            continue
        etapa = etapa_lembrete(conta, agora)
        if not etapa:
            continue
        total = len(listar_historico(nome)) if etapa == "vencida" else 0
        try:
            enviar_email(destino, *email_lembrete(etapa, conta, total))
            enviados.append((nome, etapa, conta["pago_ate"]))
        except Exception:
            log.exception("Falha ao enviar lembrete '%s' para %s", etapa, nome)
            falhas += 1

    def marcar(usuarios):
        for nome, etapa, ref in enviados:
            conta = usuarios.get(nome)
            if not conta or conta.get("pago_ate") != ref:
                continue  # renovou entretanto: o controle recomeça
            registro = conta.get("lembretes") or {}
            if registro.get("ref") != ref:
                registro = {"ref": ref, "enviados": []}
            if etapa not in registro["enviados"]:
                registro["enviados"].append(etapa)
            conta["lembretes"] = registro
        for nome, etapa, ref in enviados_teste:
            conta = usuarios.get(nome)
            if not conta:
                continue
            registro = conta.get("lembretes_teste") or {}
            if registro.get("ref") != ref:
                registro = {"ref": ref, "enviados": []}
            if etapa not in registro["enviados"]:
                registro["enviados"].append(etapa)
            conta["lembretes_teste"] = registro

    if enviados or enviados_teste:
        try:
            alterar_usuarios(marcar)
        except Exception:
            log.exception("Lembretes enviados mas não marcados; podem repetir amanhã")
    log.info("Lembretes: %s enviados, %s avisos de teste, %s falhas", len(enviados), len(enviados_teste), falhas)
    return jsonify(enviados=len(enviados), avisos_teste=len(enviados_teste), falhas=falhas)


@app.context_processor
def dados_assinatura():
    """Aviso de vencimento próximo, mostrado no topo das páginas."""
    if not session.get("assinante") or not session.get("pago_ate"):
        return {}
    restante = session["pago_ate"] - time.time()
    if 0 < restante <= AVISO_VENCIMENTO_DIAS * DIA:
        return {"dias_para_vencer": max(1, round(restante / DIA))}
    return {}


# ---------- Faixa de divulgação nos quizzes compartilhados ----------
# O quiz só traz <script src=".../divulgacao.js">: quem decide se a faixa aparece
# (e com que texto) é o servidor, na hora em que o aluno abre. Desligar no /admin
# vale para todos os quizzes, inclusive os já compartilhados.
ARQUIVO_DIVULGACAO = "config/divulgacao.json"
DIVULGACAO_PADRAO = {"ligada": True, "texto": "✨ Gere você mesmo seus quizzes a partir de textos ou fotos.",
                     "destaque": "Custa menos de R$ 1,00 por dia!", "botao": "Quero criar os meus ➜"}
_divulgacao_cache = {"dados": None, "quando": 0}


def ler_divulgacao(fresco=False):
    if not fresco and _divulgacao_cache["dados"] and time.time() - _divulgacao_cache["quando"] < 60:
        return _divulgacao_cache["dados"]
    try:
        dados, _ = ler_json_versionado(ARQUIVO_DIVULGACAO)
    except Exception:
        log.exception("Falha ao ler a configuração da divulgação")
        dados = None
    dados = {**DIVULGACAO_PADRAO, **(dados or {})}
    _divulgacao_cache.update(dados=dados, quando=time.time())
    return dados


FAIXA_JS = """(function () {
  var c = %s;
  function montar() {
    var st = document.createElement('style');
    st.textContent = '.qz-faixa{position:fixed;top:0;left:0;right:0;z-index:30;padding:10px 44px 10px 14px;background:#5920AB;color:#fff;'
      + 'font:600 13.5px/1.35 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;box-shadow:0 4px 14px rgba(0,0,0,.35);text-align:left}'
      + '.qz-faixa b{color:#C1FC38}.qz-faixa a{display:inline-block;margin-top:6px;background:#C1FC38;color:#24133E;font-weight:800;'
      + 'text-decoration:none;padding:6px 12px;border-radius:999px}.qz-faixa button{position:absolute;top:6px;right:8px;background:none;'
      + 'border:0;color:#E2D7F5;font-size:18px;padding:4px 8px;cursor:pointer}';
    document.head.appendChild(st);
    var f = document.createElement('div'); f.className = 'qz-faixa';
    var t = document.createElement('span'); t.textContent = c.texto + ' ';
    var d = document.createElement('b'); d.textContent = c.destaque;
    var a = document.createElement('a'); a.href = c.link; a.target = '_blank'; a.rel = 'noopener'; a.textContent = c.botao;
    var x = document.createElement('button'); x.type = 'button'; x.setAttribute('aria-label', 'Fechar'); x.textContent = '\\u2715';
    f.appendChild(t); f.appendChild(d); f.appendChild(document.createElement('br')); f.appendChild(a); f.appendChild(x);
    document.body.appendChild(f);
    var antes = parseFloat(getComputedStyle(document.body).paddingTop) || 0;
    function ajustar() { document.body.style.paddingTop = (f.style.display === 'none' ? antes : antes + f.offsetHeight + 8) + 'px'; }
    function mostrar() { f.style.display = ''; ajustar(); }
    ajustar(); window.addEventListener('resize', ajustar);
    // fechar esconde só até o aluno refazer o quiz (ou recomeçar os flashcards)
    x.onclick = function () { f.style.display = 'none'; ajustar(); };
    document.addEventListener('click', function (e) {
      if (e.target.closest && e.target.closest('#btn-tudo, #btn-erradas, #btn-recomecar')) mostrar();
    });
  }
  if (document.body) montar(); else document.addEventListener('DOMContentLoaded', montar);
})();"""


@app.get("/divulgacao.js")
def divulgacao_js():
    cfg = ler_divulgacao()
    tipo = request.args.get("tipo") if request.args.get("tipo") in ("quiz", "flashcards") else "quiz"
    if cfg.get("ligada"):
        dados = {k: str(cfg.get(k) or "")[:160] for k in ("texto", "destaque", "botao")}
        dados["link"] = f"{URL_APP}/assine?origem={tipo}"
        corpo = FAIXA_JS % json.dumps(dados, ensure_ascii=False).replace("</", "<\\/")
    else:
        corpo = "/* faixa desligada */"
    resposta = Response(corpo, mimetype="application/javascript")
    resposta.headers["Cache-Control"] = "public, max-age=300"   # desligar vale em até 5 minutos
    return resposta


# ---------- Teste grátis: aviso no topo e e-mail para os avisos ----------
@app.context_processor
def dados_teste():
    expira, usuario = session.get("expira_em"), session.get("usuario")
    if not expira or not usuario or session.get("assinante") or request.endpoint in ("login", "cadastro"):
        return {}
    restante = expira - time.time()
    if restante <= 0:
        return {}
    if restante < DIA:
        texto = "hoje" if datetime.fromtimestamp(expira, FUSO_BRASIL).date() == datetime.now(FUSO_BRASIL).date() else "amanhã"
    else:
        dias = int(restante // DIA) + 1
        texto = f"em {dias} dias"
    return {"teste_aviso": texto, "teste_pedir_email": not session.get("tem_email") and "@" not in usuario}


@app.post("/meu-email")
@login_obrigatorio
def meu_email():
    email = request.form.get("email", "").strip().lower()
    if not EMAIL_VALIDO.fullmatch(email) or len(email) > 120:
        flash("Informe um e-mail válido.", "erro")
        return redirect(request.referrer or url_for("index"))
    nome = session["usuario"]

    def salvar(usuarios):
        conta = usuarios.get(nome)
        if not conta:
            return False
        conta["email"] = email
        return conta
    try:
        alterar_usuarios(salvar)
    except Exception:
        log.exception("Falha ao salvar e-mail de %s", nome)
        flash("Não foi possível salvar agora. Tente novamente.", "erro")
        return redirect(request.referrer or url_for("index"))
    session["tem_email"] = True
    flash("E-mail salvo! Vamos avisar antes do teste acabar. 📬", "ok")
    destino = request.referrer or ""
    return redirect(destino if destino.startswith(request.host_url) else url_for("index"))


# ---------- Instalar o app: abrir no navegador do celular já logado ----------
# Dentro do Instagram/Facebook não dá para instalar. O aviso mostra um botão que
# abre o Chrome/Safari num link de uso rápido (10 min) que já entra na conta.
VALIDADE_LINK_NAVEGADOR = 600
SEM_AVISO_INSTALAR = {"login", "cadastro", "trocar_senha", "assinatura_pagar", "assinatura_retorno",
                      "esqueci_senha", "nova_senha"}


def _assinador_navegador():
    return URLSafeTimedSerializer(app.secret_key, salt="abrir-no-navegador")


@app.context_processor
def dados_instalar():
    """Liga o aviso de instalar só para quem já está usando o app (e, na página
    de assinatura, só depois de pago)."""
    usuario = session.get("usuario")
    if not usuario or request.endpoint in SEM_AVISO_INSTALAR or session.get("trocar_senha"):
        return {}
    if request.endpoint == "assinatura" and session.get("assinante") and (session.get("pago_ate") or 0) < time.time():
        return {}
    token = _assinador_navegador().dumps({"u": usuario})
    return {"instalar_link": url_publica("entrar_navegador", token=token)}


@app.get("/entrar/<token>")
def entrar_navegador(token):
    try:
        dados = _assinador_navegador().loads(token, max_age=VALIDADE_LINK_NAVEGADOR)
    except BadSignature:
        flash("Este link expirou. Entre com o seu e-mail e senha.", "erro")
        return redirect(url_for("login"))
    nome = dados.get("u", "")
    conta = carregar_usuarios().get(nome)
    if not conta or expirou(conta):
        return redirect(url_for("login"))
    entrar_na_sessao(nome, conta)
    return redirect(url_for("index", instalar=1))


@app.get("/")
@login_obrigatorio
def index():
    return render_template("index.html", max_arquivos=MAX_ARQUIVOS, max_mb=MAX_TOTAL_MB,
                           min_itens=MIN_ITENS, max_itens=MAX_ITENS,
                           recentes=listar_historico(session["usuario"])[:5])


def assinatura_vencida_agora():
    """Confere no Firebase (não só na sessão) se um assinante está em dia."""
    conta = carregar_usuarios().get(session["usuario"])
    if conta:
        atualizar_sessao_assinatura(conta)
    return bool(conta) and not assinatura_ativa(conta)


@app.get("/meus-quizzes")
@login_obrigatorio
def meus_quizzes():
    if session.get("assinante") and assinatura_vencida_agora():
        return redirect(url_for("assinatura"))
    return render_template("meus_quizzes.html", itens=listar_historico(session["usuario"]))


def _item_do_usuario(item_id):
    if session.get("assinante") and assinatura_vencida_agora():
        abort(402)
    item = next((i for i in listar_historico(session["usuario"]) if i.get("id") == item_id), None)
    if item is None:
        abort(404)
    return item


@app.get("/meus-quizzes/<item_id>/baixar")
@login_obrigatorio
def baixar_quiz(item_id):
    item = _item_do_usuario(item_id)
    try:
        conteudo = bucket_firebase().blob(item["arquivo"]).download_as_bytes()
    except Exception:
        log.exception("Falha ao baixar %s", item.get("arquivo"))
        flash("Não foi possível baixar este quiz agora. Tente novamente.", "erro")
        return redirect(url_for("meus_quizzes"))
    return Response(conteudo, mimetype="text/html; charset=utf-8", headers={
        "Content-Disposition": f'attachment; filename="{nome_seguro(item["titulo"])}.html"'})


@app.post("/meus-quizzes/<item_id>/remover")
@login_obrigatorio
def remover_quiz(item_id):
    if session.get("assinante") and assinatura_vencida_agora():
        return redirect(url_for("assinatura"))

    def remover(itens):
        restantes = [i for i in itens if i.get("id") != item_id]
        if len(restantes) == len(itens):
            return False  # já não estava na lista
        itens[:] = restantes
    try:
        removido = alterar_historico(session["usuario"], remover)
    except Exception:
        log.exception("Falha ao remover do histórico")
        removido = None
    if removido:
        flash("Quiz removido da lista. O link continua funcionando para quem já o recebeu.", "ok")
    elif removido is None:
        flash("Não foi possível remover agora. Tente novamente.", "erro")
    return redirect(url_for("meus_quizzes"))


# --------------------------------------------------------------------------- #
# Demonstração para a landing page (quizzify.com.br/#demo): o mesmo quiz.html
# que os professores recebem, com 5 perguntas fixas, aberto sem login.
# --------------------------------------------------------------------------- #
LANDING_URL = os.environ.get("LANDING_URL", "https://quizzify.com.br").rstrip("/")
QUIZ_DEMO = [
    {"pergunta": "Qual organela da célula é responsável pela respiração celular, que produz a maior parte da energia (ATP)?",
     "opcoes": ["Ribossomo", "Mitocôndria", "Complexo golgiense", "Lisossomo"], "correta": 1,
     "explicacao": "A mitocôndria realiza a respiração celular aeróbica, produzindo a maior parte do ATP usado pela célula."},
    {"pergunta": "Em que ano foi promulgada a Constituição Federal brasileira atual, conhecida como \"Constituição Cidadã\"?",
     "opcoes": ["1946", "1967", "1988", "1991"], "correta": 2,
     "explicacao": "A Constituição atual foi promulgada em 5 de outubro de 1988, após o fim da ditadura militar."},
    {"pergunta": "Qual gás é o principal responsável pela intensificação do efeito estufa causada pela atividade humana?",
     "opcoes": ["Oxigênio (O₂)", "Nitrogênio (N₂)", "Dióxido de carbono (CO₂)", "Hélio (He)"], "correta": 2,
     "explicacao": "A queima de combustíveis fósseis e o desmatamento aumentam a concentração de CO₂, que retém calor na atmosfera."},
    {"pergunta": "Um produto custa R$ 80,00 e está com 15% de desconto. Qual é o preço final?",
     "opcoes": ["R$ 65,00", "R$ 68,00", "R$ 70,00", "R$ 72,00"], "correta": 1,
     "explicacao": "15% de 80 é 12 (0,15 × 80). Então 80 − 12 = R$ 68,00."},
    {"pergunta": "Quem é o autor do romance \"Dom Casmurro\", marco do Realismo brasileiro?",
     "opcoes": ["José de Alencar", "Machado de Assis", "Aluísio Azevedo", "Graciliano Ramos"], "correta": 1,
     "explicacao": "\"Dom Casmurro\" (1899), de Machado de Assis, é narrado por Bentinho e famoso pela dúvida sobre Capitu."},
]


@app.get("/demo")
def demo():
    itens = [{"id": i, **q} for i, q in enumerate(QUIZ_DEMO)]
    # Na landing /assine o botão final fecha o teste e destaca os planos da própria tela;
    # na landing principal leva à seção de preços.
    na_assine = request.args.get("destino") == "assine"
    resposta = Response(render_template("quiz.html", titulo="Mini-quiz ENEM", professor="", itens=itens,
                                        cta_link=url_for("assine") if na_assine else LANDING_URL + "/#precos",
                                        cta_avisar=na_assine))
    resposta.headers["Cache-Control"] = "public, max-age=600"
    return resposta


QUIZ_VIDEO = [
    {"pergunta": "Qual organela produz a maior parte da energia da célula?",
     "opcoes": ["Ribossomo", "Mitocôndria", "Lisossomo", "Complexo golgiense"], "correta": 1,
     "explicacao": "A mitocôndria faz a respiração celular e produz o ATP."},
    {"pergunta": "Um produto de R$ 80,00 tem 15% de desconto. Qual é o preço final?",
     "opcoes": ["R$ 65,00", "R$ 68,00", "R$ 70,00", "R$ 72,00"], "correta": 1,
     "explicacao": "15% de 80 é 12. Então 80 − 12 = R$ 68,00."},
]


# Muda sempre que o quiz ou o "vídeo" mudam: o navegador não reaproveita uma versão antiga.
VERSAO_VIDEO = hashlib.sha256(b"".join(
    open(os.path.join(app.root_path, "templates", nome), "rb").read()
    for nome in ("quiz.html", "_quiz_autoplay.html"))).hexdigest()[:8]


@app.get("/demo/video")
def demo_video():
    """O quiz real jogando sozinho, em loop, dentro do celular da landing /assine."""
    itens = [{"id": i, **q} for i, q in enumerate(QUIZ_VIDEO)]
    resposta = Response(render_template("quiz.html", titulo="Revisão para a prova", professor="Ana",
                                        itens=itens, autoplay=not request.args.get("parado")))
    resposta.headers["Cache-Control"] = "public, max-age=600"
    return resposta


@app.get("/assine")
def assine():
    """Landing de uma tela só (anúncios): planos anual e mensal."""
    economia = PLANOS["mensal"]["centavos"] * 12 - PLANOS["anual"]["centavos"]
    resposta = Response(render_template("assine.html", planos=PLANOS, economia=economia,
                                        versao_video=VERSAO_VIDEO, url_canonica=url_publica("assine")))
    resposta.headers["Cache-Control"] = "public, max-age=600"
    return resposta


@app.get("/assinar/<plano_id>")
def assinar_plano(plano_id):
    """Botões "Assinar agora": logado vai direto ao checkout do plano;
    sem conta, cria a conta com o plano escolhido e depois segue para o checkout."""
    if plano_id not in PLANOS:
        abort(404)
    if not session.get("usuario"):
        return redirect(url_for("cadastro", plano=plano_id))
    if not session.get("assinante") or not INFINITEPAY_HANDLE:
        return redirect(url_for("assinatura", plano=plano_id))
    try:
        return redirect(iniciar_checkout(plano_id))
    except Exception:
        log.exception("Falha ao criar checkout pela landing")
        flash("Não foi possível abrir o pagamento agora. Tente novamente em instantes.", "erro")
        return redirect(url_for("assinatura", plano=plano_id))


# --------------------------------------------------------------------------- #
# App instalável (PWA): ícone na tela inicial do celular, abre em tela cheia.
# O manifesto e o service worker ficam na raiz para valerem para o site todo.
# --------------------------------------------------------------------------- #
COR_APP = "#160B29"   # fundo do app (roxo-escuro da marca)


@app.get("/manifest.webmanifest")
def manifesto():
    icone = lambda nome, tam, uso="any": {"src": url_for("static", filename=f"icons/{nome}"),
                                         "sizes": f"{tam}x{tam}", "type": "image/png", "purpose": uso}
    dados = {
        "id": "/", "name": "Quizzify", "short_name": "Quizzify", "lang": "pt-BR",
        "description": "Transforme documentos e fotos em quizzes e flashcards com IA.",
        "start_url": "/?origem=app", "scope": "/", "display": "standalone", "orientation": "portrait",
        "background_color": COR_APP, "theme_color": COR_APP,
        "icons": [icone("icon-192.png", 192), icone("icon-512.png", 512), icone("icon-512.png", 512, "maskable")],
    }
    resposta = Response(json.dumps(dados, ensure_ascii=False), mimetype="application/manifest+json")
    resposta.headers["Cache-Control"] = "public, max-age=3600"
    return resposta


@app.get("/sw.js")
def service_worker():
    resposta = Response(render_template("sw.js"), mimetype="application/javascript")
    resposta.headers.update({"Cache-Control": "no-cache", "Service-Worker-Allowed": "/"})
    return resposta


@app.after_request
def cache_arquivos_fixos(resposta):
    """Fontes e ícones quase nunca mudam: o navegador (e a Cloudflare do Render)
    guardam-nos, e as páginas abrem mais rápido na visita seguinte."""
    if request.path.startswith("/static/fonts/"):
        resposta.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    elif request.path.startswith("/static/icons/"):
        resposta.headers["Cache-Control"] = "public, max-age=604800"
    return resposta


# --------------------------------------------------------------------------- #
# Google: só a página de vendas (/assine) aparece nas buscas. Login, cadastro,
# admin, quizzes e o resto do app ficam fora (noindex), para o resultado do
# Google mostrar a página certa.
# --------------------------------------------------------------------------- #
PAGINAS_PUBLICAS = {"assine", "robots", "sitemap"}


@app.get("/robots.txt")
def robots():
    texto = ("User-agent: *\n"
             "Allow: /assine\n"
             "Disallow: /admin\nDisallow: /api/\nDisallow: /meus-quizzes\nDisallow: /tarefas/\n"
             "Disallow: /webhooks/\nDisallow: /m/\nDisallow: /assinar/\n"
             f"\nSitemap: {url_publica('sitemap')}\n")
    return Response(texto, mimetype="text/plain")


@app.get("/sitemap.xml")
def sitemap():
    xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
           f'  <url><loc>{url_publica("assine")}</loc><changefreq>weekly</changefreq><priority>1.0</priority></url>\n'
           '</urlset>\n')
    return Response(xml, mimetype="application/xml")


@app.after_request
def fora_do_google(resposta):
    if request.endpoint not in PAGINAS_PUBLICAS and not request.path.startswith("/static/"):
        resposta.headers["X-Robots-Tag"] = "noindex"
    return resposta


@app.get("/health")
def health():
    """Rota leve para o cron-job manter o serviço acordado no Render."""
    return "ok", 200


@app.post("/api/gerar")
@login_obrigatorio
def gerar():
    try:
        usuarios = carregar_usuarios(estrito=True)
    except FalhaUsuarios:
        return jsonify(erro="Serviço temporariamente indisponível. Tente novamente."), 503
    if not conta_ainda_existe(usuarios):
        return jsonify(erro="A sua conta já não está ativa. Entre novamente."), 401
    usuario = session["usuario"]
    conta = usuarios[usuario]
    if not assinatura_ativa(conta):
        atualizar_sessao_assinatura(conta)
        return jsonify(erro="A sua assinatura não está ativa. Renove para continuar."), 402
    sem_limite = bool(conta.get("admin"))
    limite = LIMITE_TESTE if conta.get("expira_em") else LIMITE_DIARIO

    arquivos = [a for a in request.files.getlist("arquivos") if a and a.filename]
    tipo = request.form.get("tipo", "quiz")
    titulo = (request.form.get("titulo") or "Desafio de Conhecimento").strip()[:120]
    professor = (request.form.get("professor") or "").strip()[:80]
    divulgar = request.form.get("divulgar") == "1"
    try:
        n = min(max(int(request.form.get("quantidade", 10)), MIN_ITENS), MAX_ITENS)
    except ValueError:
        n = 10

    if tipo not in ("quiz", "flashcards"):
        return jsonify(erro="Tipo inválido."), 400
    if not arquivos:
        return jsonify(erro="Envie pelo menos um arquivo ou foto."), 400
    if len(arquivos) > MAX_ARQUIVOS:
        return jsonify(erro=f"Envie no máximo {MAX_ARQUIVOS} arquivos de cada vez."), 400
    invalidos = [a.filename for a in arquivos
                 if extensao(a.filename) not in EXTENSOES_DOC | EXTENSOES_IMG
                 and not (a.mimetype or "").startswith("image/")]
    if invalidos:
        return jsonify(erro=f"Formato não suportado: {', '.join(invalidos)}"), 400
    with vaga_geracao() as tem_vaga:
        if not tem_vaga:
            return jsonify(erro="Muitas pessoas gerando quizzes agora. Tente de novo em 1 minuto."), 503
        if not sem_limite and (aviso := verificar_limite(usuario, limite)):
            return jsonify(erro=aviso), 429

        try:
            texto, partes = preparar_conteudo(arquivos)
        except Exception:
            log.exception("Falha na extração")
            return jsonify(erro="Não foi possível ler um dos arquivos. Verifique se não está corrompido."), 422
        if not texto and not partes:
            return jsonify(erro="Não encontrei texto nos arquivos enviados."), 422

        try:
            if tipo == "quiz":
                itens = normalizar_quiz(chamar_gemini(PROMPT_QUIZ.format(n=n), texto, partes), n)
            else:
                itens = normalizar_flash(chamar_gemini(PROMPT_FLASH.format(n=n), texto, partes), n)
        except Exception:
            log.exception("Falha na geração com IA")
            return jsonify(erro="A IA não conseguiu gerar o conteúdo. Tente novamente."), 502
        if not itens:
            return jsonify(erro="A IA devolveu um resultado vazio. Tente novamente."), 502

        template = "quiz.html" if tipo == "quiz" else "flashcards.html"
        conteudo_html = render_template(template, titulo=titulo, professor=professor, itens=itens,
                                        divulgacao_js=f"{URL_APP}/divulgacao.js?tipo={tipo}" if divulgar else None)
        link, caminho = publicar_firebase(conteudo_html, tipo)
        if not sem_limite:
            registrar_geracao(usuario)
        item = None
        if link:
            item = {"id": uuid.uuid4().hex[:12], "titulo": titulo, "tipo": tipo, "total": len(itens),
                    "link": link, "arquivo": caminho, "criado_em": time.time()}
            registrar_no_historico(usuario, item)

        return jsonify(html=conteudo_html, link=link, total=len(itens),
                       nome_arquivo=f"{nome_seguro(titulo)}.html",
                       item=item and {**item, "data": data_hora(item["criado_em"])})


@app.errorhandler(RequestEntityTooLarge)
def arquivo_grande(_):
    return jsonify(erro=f"O tamanho total dos arquivos excede {MAX_TOTAL_MB} MB."), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
